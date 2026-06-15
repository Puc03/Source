from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs" / "reports"
OUTPUT_PATH = OUTPUT_DIR / "BIG_DATA_Nhom_8_Education_enterprise_first_revised.docx"


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(9)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "B7C7D9")
        borders.append(elem)
    tbl_pr.append(borders)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True)
        shade_cell(hdr[i], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table)
    cell = table.rows[0].cells[0]
    shade_cell(cell, "EAF4E4")
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(49, 94, 37)
    b = p.add_run(body)
    b.font.name = "Arial"
    b.font.size = Pt(10)
    doc.add_paragraph()


def setup_styles(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size, color in [
        ("Heading 1", 16, RGBColor(31, 78, 121)),
        ("Heading 2", 13, RGBColor(55, 86, 35)),
        ("Heading 3", 11.5, RGBColor(79, 79, 79)),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def build_doc() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_styles(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("UEH Smart Library Recommendation System")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Báo cáo demo Big Data theo hướng Enterprise Architecture trước, Proof of Concept sau")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.italic = True

    doc.add_paragraph("GitHub source: https://github.com/Puc03/Source.git")
    add_callout(
        doc,
        "Cách đọc báo cáo",
        "Báo cáo trình bày hệ thống đầy đủ trong phạm vi doanh nghiệp trước, sau đó mới mô tả bản demo rút gọn mà nhóm đã cài đặt bằng dữ liệu mô phỏng, DuckDB, Parquet, PySpark và Power BI.",
    )

    doc.add_heading("1. Giới Thiệu Bài Toán", level=1)
    doc.add_paragraph(
        "Trong môi trường thư viện đại học, sinh viên phải tìm kiếm tài liệu từ nhiều nguồn khác nhau như sách in, tài liệu số, luận văn, cơ sở dữ liệu học thuật và các nguồn tham khảo bên ngoài. Khi số lượng tài nguyên ngày càng lớn, người học dễ gặp khó khăn trong việc chọn tài liệu phù hợp với ngành học, môn học và nhu cầu nghiên cứu."
    )
    doc.add_paragraph(
        "Nhóm đề xuất hệ thống UEH Smart Library Recommendation System nhằm phân tích dữ liệu thư viện, dữ liệu học tập và dữ liệu tương tác người dùng để gợi ý tài nguyên phù hợp cho từng sinh viên. Trong phạm vi doanh nghiệp, hệ thống có thể kết nối với OPAC/Sierra, DSpace Repository, EBSCO OneSearch, OpenAthens, LMS và các kênh truy cập thư viện khác."
    )
    doc.add_paragraph(
        "Do nhóm không có quyền truy cập trực tiếp vào cơ sở dữ liệu nội bộ và log hành vi thật của sinh viên, phần cài đặt được triển khai dưới dạng demo rút gọn. Demo sử dụng dữ liệu metadata công khai, proxy data từ Tiki và dữ liệu sinh viên mô phỏng để chứng minh pipeline Big Data và mô hình đề xuất có thể vận hành."
    )

    doc.add_heading("2. Kiến Trúc Hệ Thống Đầy Đủ Trong Phạm Vi Doanh Nghiệp", level=1)
    doc.add_heading("2.1. Mục Tiêu Hệ Thống Production", level=2)
    add_bullets(
        doc,
        [
            "Hợp nhất dữ liệu tài nguyên, dữ liệu học tập và dữ liệu hành vi từ nhiều hệ thống thư viện.",
            "Bảo vệ dữ liệu cá nhân sinh viên trước khi đưa vào phân tích.",
            "Xây dựng mô hình đề xuất tài liệu phù hợp theo ngành học, môn học, lịch sử tương tác và mức độ phổ biến của tài nguyên.",
            "Phân phối kết quả đề xuất đến web portal, mobile app, email digest hoặc OPAC terminal widget.",
        ],
    )

    doc.add_heading("2.2. Nguồn Dữ Liệu Trong Hệ Thống Thật", level=2)
    add_table(
        doc,
        ["Nguồn dữ liệu", "Vai trò trong hệ thống"],
        [
            ["OPAC / Sierra Circulation", "Lịch sử mượn, trả, gia hạn, đặt giữ chỗ sách in."],
            ["DSpace / UEH Repository", "Metadata tài liệu số, lượt xem, lượt tải, collection, DOI, tác giả."],
            ["EBSCO OneSearch", "Từ khóa tìm kiếm, hành vi click kết quả, mở cơ sở dữ liệu học thuật."],
            ["OpenAthens", "Log xác thực và truy cập tài nguyên điện tử."],
            ["LMS / hệ thống học tập", "Thông tin ngành, môn học, học phần và ngữ cảnh học tập."],
            ["Feedback sinh viên", "Rating, click recommendation, lưu tài liệu, bỏ qua đề xuất."],
            ["Nguồn ngoài như Tiki", "Proxy data về sách: rating, review_count, quantity_sold, giá và URL sản phẩm."],
        ],
    )

    doc.add_heading("2.3. Tầng Thu Thập Dữ Liệu", level=2)
    doc.add_paragraph(
        "Trong hệ thống production, dữ liệu có thể được thu thập theo hai cơ chế. Batch ingestion dùng để nhập dữ liệu định kỳ từ OPAC, repository, LMS hoặc API. Streaming ingestion dùng Kafka để nhận sự kiện gần thời gian thực từ web portal, mobile app, OPAC terminal hoặc hệ thống xác thực."
    )
    add_table(
        doc,
        ["Loại sự kiện", "Ý nghĩa"],
        [
            ["student_search", "Sinh viên tìm kiếm tài liệu."],
            ["item_view", "Sinh viên xem chi tiết một tài nguyên."],
            ["repository_download", "Sinh viên tải hoặc mở file PDF từ repository."],
            ["book_borrow / book_return", "Sinh viên mượn hoặc trả sách in."],
            ["hold_request", "Sinh viên đặt giữ chỗ tài liệu."],
            ["recommendation_click", "Sinh viên click vào tài nguyên được đề xuất."],
        ],
    )

    doc.add_heading("2.4. Privacy Layer Và Data Governance", level=2)
    doc.add_paragraph(
        "Vì dữ liệu sinh viên là dữ liệu nhạy cảm, hệ thống cần có privacy layer trước khi đưa dữ liệu vào phân tích. Lớp này tách thông tin định danh khỏi dữ liệu analytics và thay bằng token ẩn danh."
    )
    add_bullets(
        doc,
        [
            "Sinh student_token để thay thế mã sinh viên thật.",
            "Hash hoặc mã hóa định danh nhạy cảm.",
            "Loại bỏ các trường như student_id, student_code, student_name, email khỏi dữ liệu analytics.",
            "Phân quyền dữ liệu theo vai trò người dùng.",
            "Theo dõi lineage bằng công cụ như Apache Atlas hoặc DataHub nếu triển khai production.",
        ],
    )

    doc.add_heading("2.5. Data Lake Và Lưu Trữ Phân Tán", level=2)
    doc.add_paragraph(
        "Ở quy mô doanh nghiệp, dữ liệu nên được lưu trong Data Lake trên HDFS, S3 hoặc MinIO. Định dạng Parquet phù hợp cho dữ liệu phân tích vì hỗ trợ nén tốt, đọc theo cột và tối ưu cho truy vấn OLAP."
    )
    add_table(
        doc,
        ["Tầng dữ liệu", "Nội dung"],
        [
            ["Bronze", "Dữ liệu thô từ OPAC, Repository, OneSearch, OpenAthens, LMS."],
            ["Silver", "Dữ liệu đã làm sạch, chuẩn hóa, ẩn danh."],
            ["Gold", "Feature cho recommendation, bảng phân tích, kết quả đề xuất."],
        ],
    )

    doc.add_heading("2.6. Processing, Warehouse Và AI Recommendation", level=2)
    doc.add_paragraph(
        "Trong production, Spark có thể xử lý dữ liệu lớn, tạo feature và huấn luyện mô hình. Sau đó các bảng tổng hợp có thể đưa vào Hive, ClickHouse hoặc một warehouse phân tích để dashboard truy vấn nhanh."
    )
    add_table(
        doc,
        ["Mô hình", "Vai trò"],
        [
            ["Content-based Recommendation", "Đề xuất dựa trên ngành học, môn học, tiêu đề, tác giả, chủ đề và metadata tài nguyên."],
            ["Collaborative Filtering ALS", "Đề xuất dựa trên hành vi tương tác thật của nhiều sinh viên."],
            ["Semantic Search / Embedding", "Đo độ tương đồng ngữ nghĩa giữa nhu cầu học tập và nội dung tài liệu."],
            ["Hybrid Ranking", "Kết hợp match score, hành vi, rating, độ phổ biến, độ mới và quyền truy cập."],
        ],
    )

    doc.add_heading("2.7. Serving Và Phân Phối Đề Xuất", level=2)
    doc.add_paragraph(
        "Sau khi mô hình tạo danh sách đề xuất, hệ thống cần kiểm tra tính khả dụng của tài nguyên. Sách in cần kiểm tra trạng thái còn/hết và vị trí kệ; tài liệu số cần kiểm tra quyền truy cập, link PDF, DOI hoặc repository URL. Kết quả có thể được phân phối qua web portal, mobile app, email digest hoặc OPAC terminal widget."
    )

    doc.add_heading("3. Bản Demo Rút Gọn Của Nhóm", level=1)
    add_callout(
        doc,
        "Vai trò của demo",
        "Demo không phải hệ thống production hoàn chỉnh. Demo là proof of concept chạy local để chứng minh pipeline Big Data và mô hình đề xuất có thể hoạt động trong điều kiện không có log thật từ hệ thống nội bộ.",
    )

    doc.add_heading("3.1. Dữ Liệu Demo", level=2)
    add_table(
        doc,
        ["Nhóm dữ liệu", "Mục đích sử dụng"],
        [
            ["UEH Repository metadata", "Đại diện cho tài liệu số học thuật của UEH."],
            ["Tiki books proxy data", "Bổ sung tín hiệu rating, review_count, quantity_sold, giá và URL cho sách tham khảo."],
            ["Mock students", "Mô phỏng 2,000 sinh viên theo ngành học, khóa học và môn học."],
            ["Simulated behavior events", "Mô phỏng hành vi từ OpenAthens, DSpace, EBSCO OneSearch và Sierra."],
        ],
    )
    doc.add_paragraph(
        "Dữ liệu Tiki chỉ là proxy data, không phải dữ liệu hành vi thật của sinh viên UEH. Các trường rating_average, review_count và quantity_sold_value được dùng để mô phỏng tín hiệu chất lượng và độ phổ biến của sách."
    )

    doc.add_heading("3.2. Privacy Layer Trong Demo", level=2)
    doc.add_paragraph(
        "Nhóm tạo student_token và student_hash để thay thế thông tin định danh sinh viên. Các file analytics không chứa mã sinh viên gốc, tên sinh viên hoặc user_id."
    )
    add_table(
        doc,
        ["File / kết quả", "Ý nghĩa"],
        [
            ["student_token_map_protected.csv", "Bảng ánh xạ token đã bảo vệ, không chứa PII thô."],
            ["students_protected.csv", "Thông tin sinh viên đã ẩn danh."],
            ["student_resource_matches_analytics.csv", "Dữ liệu matching dùng cho phân tích và recommendation."],
            ["Privacy check", "Kiểm tra không còn cột PII bị cấm trong bảng analytics."],
        ],
    )

    doc.add_heading("3.3. Data Lake Local Bằng Parquet", level=2)
    add_table(
        doc,
        ["Thuộc tính", "Giá trị demo"],
        [
            ["Định dạng lưu trữ", "Parquet"],
            ["Partition key", "major"],
            ["Số dòng", "12,000"],
            ["Số partition", "36"],
            ["Vai trò", "Mô phỏng cách tổ chức dữ liệu trên HDFS/S3/MinIO."],
        ],
    )

    doc.add_heading("3.4. DuckDB Warehouse Và Analysis", level=2)
    doc.add_paragraph(
        "DuckDB được dùng làm analytical warehouse local để đọc dữ liệu Parquet, tạo bảng fact và thực hiện truy vấn OLAP. Công cụ này phù hợp cho demo vì nhẹ, chạy cục bộ và vẫn hỗ trợ truy vấn phân tích mạnh."
    )
    add_table(
        doc,
        ["KPI", "Giá trị"],
        [
            ["Total match rows", "12,000"],
            ["Total students", "2,000"],
            ["Total majors", "36"],
            ["Total courses", "201"],
            ["Tiki books", "162"],
            ["UEH repository documents", "152"],
            ["Average Tiki match score", "50.91"],
            ["Average repository match score", "40.44"],
        ],
    )

    doc.add_heading("3.5. AI Recommendation Trong Demo", level=2)
    doc.add_paragraph(
        "Mô hình chính trong demo là content-based recommendation kết hợp weighted scoring. Điểm đề xuất được tính từ match score, confidence, rating Tiki, review_count, quantity_sold, metadata repository và độ phù hợp với ngành học/môn học."
    )
    add_table(
        doc,
        ["Chỉ số recommendation", "Giá trị"],
        [
            ["Top-5 recommendation rows", "10,000"],
            ["Students with recommendations", "2,000"],
            ["Recommended resources", "151"],
            ["Average recommendation score", "70.5"],
            ["Minimum recommendation score", "31.28"],
            ["Maximum recommendation score", "96.42"],
        ],
    )

    doc.add_heading("3.6. PySpark ALS Demo Mở Rộng", level=2)
    doc.add_paragraph(
        "Ngoài mô hình chính, nhóm triển khai PySpark ALS để minh họa hướng phát triển collaborative filtering trong môi trường Big Data. ALS trong demo dùng proxy implicit rating, chưa phải rating thật của sinh viên."
    )
    add_table(
        doc,
        ["Chỉ số Spark ALS", "Giá trị"],
        [
            ["Proxy interactions", "22,115"],
            ["Users", "2,000"],
            ["Items", "314"],
            ["RMSE", "4.7891"],
            ["Vai trò", "Demo mở rộng, không thay thế mô hình chính."],
        ],
    )

    doc.add_heading("3.7. Phase 5: Phân Phối Kết Quả Đề Xuất", level=2)
    doc.add_paragraph(
        "Phase 5 mô phỏng bước đưa kết quả đề xuất đến người dùng cuối. Kết quả được chuẩn hóa thành payload để có thể hiển thị trên web portal hoặc mobile app."
    )
    add_table(
        doc,
        ["KPI phân phối", "Giá trị"],
        [
            ["Delivery rows", "10,000"],
            ["Students", "2,000"],
            ["Resources", "151"],
            ["Tiki / print / marketplace items", "5,394"],
            ["Digital repository items", "4,606"],
            ["Average recommendation score", "70.5"],
        ],
    )

    doc.add_heading("4. Dashboard Power BI", level=1)
    doc.add_paragraph(
        "Power BI được dùng để trực quan hóa kết quả demo theo các nhóm: tổng quan KPI, hành vi mô phỏng, matching analysis, recommendation result và distribution output. Dashboard giúp người xem hiểu luồng dữ liệu từ đầu vào đến kết quả đề xuất cuối cùng."
    )
    add_numbered(
        doc,
        [
            "Overview: trình bày KPI tổng quan và pipeline.",
            "Behavior Analysis: phân tích hành vi người dùng mô phỏng từ 4 hệ thống.",
            "Matching Analysis: trình bày match rows, confidence và điểm matching.",
            "Recommendation: trình bày Top-N resources, score và tài nguyên được đề xuất.",
            "Distribution: trình bày payload và trạng thái khả dụng của tài nguyên.",
        ],
    )

    doc.add_heading("5. Insight Chính Từ Demo", level=1)
    add_bullets(
        doc,
        [
            "Hệ thống tạo được đề xuất cho toàn bộ 2,000 sinh viên trong dữ liệu mô phỏng.",
            "Các ngành tài chính, kinh doanh và công nghệ có số lượng match nổi bật trong dữ liệu demo.",
            "Tiki books bổ sung tín hiệu phổ biến thông qua rating, review_count và quantity_sold.",
            "UEH Repository đóng vai trò nguồn học thuật chính thống, bổ sung tài liệu số cho recommendation.",
            "Content-based recommendation phù hợp trong giai đoạn chưa có log hành vi thật.",
            "PySpark ALS là hướng phát triển hợp lý khi hệ thống có đủ lịch sử mượn/trả, click, tải PDF hoặc feedback thật.",
        ],
    )

    doc.add_heading("6. Hạn Chế", level=1)
    add_bullets(
        doc,
        [
            "Dữ liệu hành vi sinh viên là mô phỏng, chưa phải log thật từ UEH Smart Library.",
            "Dữ liệu Tiki là proxy data, không phải hành vi thật của sinh viên UEH.",
            "Data Lake đang chạy local bằng Parquet, chưa phải HDFS/S3/MinIO thật.",
            "DuckDB là analytical warehouse local, chưa phải warehouse production.",
            "Phase 5 chưa kết nối OPAC circulation API hoặc repository access service real-time.",
            "PySpark ALS dùng proxy implicit rating, chưa dùng rating thật.",
        ],
    )

    doc.add_heading("7. Hướng Phát Triển", level=1)
    add_bullets(
        doc,
        [
            "Thu thập log thật từ OPAC, Repository, LMS, OpenAthens và web portal.",
            "Lưu trữ dữ liệu trên HDFS, S3 hoặc MinIO.",
            "Dùng Spark để xử lý dữ liệu lớn và huấn luyện recommendation model.",
            "Triển khai Spark ALS trên dữ liệu hành vi thật.",
            "Bổ sung semantic search bằng embedding và vector database.",
            "Xây API Gateway để phân phối đề xuất tới web portal, mobile app và email digest.",
            "Đánh giá chất lượng đề xuất bằng click-through rate, borrow rate, download rate và feedback thật.",
        ],
    )

    doc.add_heading("8. Kết Luận", level=1)
    doc.add_paragraph(
        "Đề tài đã xây dựng được một bản demo hoàn chỉnh cho hệ thống đề xuất tài nguyên thư viện thông minh. Dù chưa sử dụng dữ liệu hành vi thật và chưa triển khai trên hạ tầng doanh nghiệp, demo đã thể hiện đầy đủ các thành phần quan trọng của một pipeline Big Data: thu thập dữ liệu, bảo vệ quyền riêng tư, lưu trữ dạng Data Lake, phân tích bằng warehouse, xây dựng mô hình đề xuất và phân phối kết quả."
    )
    doc.add_paragraph(
        "Hệ thống demo là nền tảng để phát triển lên phiên bản production trong tương lai, khi có thể kết nối trực tiếp với các hệ thống thật của UEH Smart Library và thu thập dữ liệu hành vi người dùng ở quy mô lớn."
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("Phụ Lục: Ảnh Minh Họa Nên Giữ Trong Báo Cáo", level=1)
    add_table(
        doc,
        ["Nhóm ảnh", "Nội dung nên chụp/giữ"],
        [
            ["Data collection", "Code crawl UEH Repository và Tiki books, kèm bảng dữ liệu mẫu."],
            ["Privacy layer", "Code tạo student_token/student_hash và query kiểm tra không còn PII."],
            ["Data Lake", "Query đọc Parquet, total_rows = 12,000, total_partitions = 36."],
            ["DuckDB analysis", "Overview KPI, major summary, confidence summary."],
            ["Recommendation", "Top resources, Top-5 recommendation cho một student_token."],
            ["PySpark ALS", "Command chạy script, input summary, RMSE, top recommendations."],
            ["Phase 5", "Distribution KPIs, availability summary, web_portal_payload."],
            ["Power BI", "Dashboard cuối cùng gồm Overview, Behavior, Matching, Recommendation và Distribution."],
        ],
    )

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_doc()
